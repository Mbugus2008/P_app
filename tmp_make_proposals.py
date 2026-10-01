"""Generate the Parcel System proposal documents (client + sales partner)."""
import os
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT_DIR = r"d:\Karai\Rembo\Proposals"
os.makedirs(OUT_DIR, exist_ok=True)

COMPANY = "[YOUR COMPANY]"
TODAY = date(2026, 9, 30).strftime("%d %B %Y")

BRAND = RGBColor(0x0B, 0x5A, 0x3C)


def title(doc, text, subtitle=None):
    h = doc.add_heading(text, level=0)
    for run in h.runs:
        run.font.color.rgb = BRAND
    if subtitle:
        p = doc.add_paragraph()
        r = p.add_run(subtitle)
        r.italic = True
        r.font.size = Pt(11)


def heading(doc, text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = BRAND
    return h


def bullets(doc, items, bold_prefix=True):
    for item in items:
        if isinstance(item, tuple):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item[0] + " ")
            run.bold = bold_prefix
            p.add_run(item[1])
        else:
            doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = w
    return t


# =====================================================================
# 1. CLIENT PROPOSAL
# =====================================================================
doc = Document()
title(
    doc,
    "Parcel Management System",
    "A complete digital platform for parcel and courier operations\n"
    f"Prepared for: [CLIENT COMPANY]        Prepared by: {COMPANY}        Date: {TODAY}",
)

heading(doc, "1. Executive Summary")
doc.add_paragraph(
    "We propose a complete, production-proven parcel management platform that digitises "
    "every step of a courier operation — parcel intake at the counter, batch dispatch by "
    "vehicle, receiving at the destination branch, customer SMS notification with a "
    "collection code, payment collection (cash or M-Pesa), and full posting into "
    "Microsoft Dynamics 365 Business Central for accounting and reporting."
)
doc.add_paragraph(
    "The system is already in daily commercial use, processing thousands of parcels, "
    "with agents, supervisors and management working from the same live data."
)

heading(doc, "2. The Problem We Solve")
bullets(
    doc,
    [
        ("Lost and untraceable parcels —", "no single record of who dispatched, carried or received each parcel."),
        ("Aging uncollected parcels —", "customers are never reminded, parcels sit for weeks, cash is delayed."),
        ("Cash reconciliation gaps —", "counter staff collect money that is never receipted or posted."),
        ("No customer communication —", "receivers don't know when a parcel arrives; senders keep calling."),
        ("Manual registers and Excel —", "owners have no real-time view of operations across branches."),
    ],
)

heading(doc, "3. How It Works")
steps = [
    "An agent registers the parcel on the Android app: sender, receiver, route, amount, who pays.",
    "Parcels are grouped into a batch and dispatched — the vehicle and driver are recorded.",
    "The receiving branch confirms the batch; each parcel gets a unique 5-digit collection code sent to the receiver by SMS.",
    "The receiver pays (cash or M-Pesa STK) and collects using the code; a thermal receipt is printed.",
    "Everything posts to Business Central automatically; owners see live dashboards and daily collection reports.",
]
for i, s in enumerate(steps, 1):
    doc.add_paragraph(f"{i}. {s}")

heading(doc, "4. What Is Included")
table(
    doc,
    ["Component", "Details"],
    [
        ("Android application", "Used by counter agents and receiving staff; works on low-cost phones and tablets."),
        ("Batch tracking", "Dispatch and receive batches per vehicle; discrepancies visible instantly."),
        ("Customer SMS", "Automatic notifications with collection codes and amounts due (sender ID branded to your business)."),
        ("Payments", "Cash and M-Pesa (STK push) with reconciliation per parcel, agent and branch."),
        ("Business Central integration", "Parcels, batches, users, locations, payments and timestamps posted to your BC environment."),
        ("Dashboards & reports", "Daily collections, location performance, follow-up reports, agent accountability."),
        ("Administration", "User roles, locations, vehicles, app auto-updates — centrally managed."),
    ],
    widths=None,
)

heading(doc, "5. Benefits")
bullets(
    doc,
    [
        ("Reduce losses:", "every parcel is traceable to the last hand that touched it."),
        ("Collect faster:", "SMS with collection codes and amounts due shortens collection cycles."),
        ("Plug cash leaks:", "every shilling collected is recorded against a parcel and an agent."),
        ("See everything live:", "own real-time numbers from any phone or PC — no waiting for month-end."),
        ("Low training cost:", "agents use the app confidently after a half-day session."),
    ],
)

heading(doc, "6. Commercial Terms")
table(
    doc,
    ["Item", "Terms"],
    [
        ("One-time setup & onboarding", "KSh [__________] — Business Central configuration, branding, user setup, staff training"),
        ("Service fee", "KSh 20 per parcel processed (includes app, SMS notifications, hosting, support and updates)"),
        ("Volume tiers", "Negotiable above 3,000 parcels per month"),
        ("Monthly minimum", "KSh 5,000 per month"),
        ("Billing", "Monthly statement; payment by M-Pesa Paybill or bank transfer within 7 days"),
    ],
)
doc.add_paragraph(
    "Note: SMS costs and software updates are included in the per-parcel fee — no hidden charges."
).italic = True

heading(doc, "7. Implementation Plan")
table(
    doc,
    ["Phase", "Duration", "Activities"],
    [
        ("1. Setup", "Week 1", "Business Central objects & connections configured; locations, users, vehicles and SMS branding set up."),
        ("2. Configuration & Training", "Week 2", "Fee/payment rules configured; staff trained on parcel entry, dispatch, receive and collection."),
        ("3. Pilot", "Week 3", "One branch/lane runs on the system with our team on standby."),
        ("4. Go-live", "Week 4", "Full rollout; hyper-care support for two weeks."),
    ],
)

heading(doc, "8. Support")
bullets(
    doc,
    [
        "Telephone / WhatsApp support during business hours; on-site support on request.",
        "Software updates delivered automatically to all devices.",
        "Quarterly operations review with your management team.",
    ],
)

heading(doc, "9. Why Us")
doc.add_paragraph(
    "We built and operate this platform in live parcel businesses — we understand Kenyan "
    "courier operations, Business Central, M-Pesa and what owners actually need from their numbers. "
    "You are not buying a prototype; you are joining a working system."
)

heading(doc, "10. Next Steps")
doc.add_paragraph(
    "To proceed: sign this proposal, confirm your setup scope, and we will begin configuration "
    "the following week. A formal service agreement will accompany this proposal."
)
doc.add_paragraph("")
table(
    doc,
    ["Accepted for [CLIENT COMPANY]", f"For and on behalf of {COMPANY}"],
    [("Name: ______________________", f"Name: ______________________"),
     ("Position: __________________", "Position: __________________"),
     ("Signature: _________________", "Signature: _________________"),
     ("Date: _____________________", "Date: _____________________")],
)

path1 = os.path.join(OUT_DIR, "CLIENT PROPOSAL - Parcel Management System.docx")
doc.save(path1)
print("saved:", path1)

# =====================================================================
# 2. SALES PARTNER PROGRAMME
# =====================================================================
doc = Document()
title(
    doc,
    "Sales Partner Programme",
    "Introduce courier and parcel businesses to a proven digital operations platform\n"
    f"Issued by: {COMPANY}        Date: {TODAY}",
)

heading(doc, "1. The Opportunity")
doc.add_paragraph(
    "Courier and parcel businesses in Kenya are moving from paper registers and Excel to "
    "digital operations. Our platform — already in daily commercial use — digitises parcel "
    "intake, dispatch, receiving, customer SMS, payments and Business Central accounting. "
    "You earn every time a business you bring in uses it."
)

heading(doc, "2. Two Ways to Earn")
table(
    doc,
    ["Role", "What you do", "What you earn"],
    [
        ("ACCOUNT PARTNER", "Introduce clients and actively follow up with them (check-ins, support coordination, retention).",
         "KSh 5 per parcel — recurring, every month, for every parcel your clients process"),
        ("REFERRAL", "Introduce a client; we handle everything after that.", "KSh 5,000 one-off per client that signs and activates"),
    ],
)

heading(doc, "3. Earnings Illustration (Account Partner)")
table(
    doc,
    ["Clients you manage", "Parcels per month", "Your monthly commission (KSh 5/parcel)"],
    [
        ("3 clients × ~500 parcels", "1,500", "7,500"),
        ("5 clients × ~900 parcels", "4,500", "22,500"),
        ("10 clients × ~1,000 parcels", "10,000", "50,000"),
    ],
)
doc.add_paragraph(
    "Commission is recurring: as long as your clients keep operating and their invoices are settled, "
    "you keep earning every month."
)

heading(doc, "4. How It Works")
for i, s in enumerate([
    "Sign the partner agreement and receive your partner kit (demo access, price list, presentation).",
    "Introduce a prospective client to us and register the lead with your name.",
    "We run the demo, survey and onboarding (including Business Central setup and staff training).",
    "When the client goes live, commission starts — tracked automatically per parcel.",
    "You receive a monthly statement of parcels processed and commission earned; payment is made after the client's monthly invoice clears.",
], 1):
    doc.add_paragraph(f"{i}. {s}")

heading(doc, "5. What We Provide")
bullets(
    doc,
    [
        "Product training and a live demo account.",
        "Ready-made presentation, price list and FAQ for your client meetings.",
        "A technical onboarding team — you don't need to know Business Central.",
        "Monthly commission statements with full parcel-level transparency.",
    ],
)

heading(doc, "6. Partner Requirements")
bullets(
    doc,
    [
        "Professional conduct and confidentiality of client information.",
        "Account Partners: a monthly check-in with each client you manage (retention is part of the job).",
        "Referrers: simply register the lead before introduction.",
    ],
)

heading(doc, "7. Programme Terms (Summary)")
table(
    doc,
    ["Item", "Term"],
    [
        ("Account Partner commission", "KSh 5 per parcel processed monthly by attributed clients"),
        ("Referral fee", "KSh 5,000 one-off, paid after the client's first invoice is settled"),
        ("Payment basis", "Commission is paid only on revenue actually collected from the client"),
        ("Statements", "Monthly, from the platform's parcel counts"),
        ("Attribution", "First registered lead wins; disputes resolved by management"),
        ("Non-solicitation", "Partners may not divert platform clients to competing services"),
    ],
)

heading(doc, "8. Join Us")
doc.add_paragraph(
    "Fill in your details below and return this page to [__________________] to be enrolled."
)
table(
    doc,
    ["Partner details", ""],
    [("Full name: ______________________________", "Phone: __________________________"),
     ("ID / Company: ___________________________", "Email: ___________________________"),
     ("I am joining as:  ☐ Account Partner      ☐ Referral", "Signature: ______________________"),
     ("KRA PIN: _______________________________", "Date: ___________________________")],
)

path2 = os.path.join(OUT_DIR, "SALES PARTNER PROGRAMME.docx")
doc.save(path2)
print("saved:", path2)
print("done")
